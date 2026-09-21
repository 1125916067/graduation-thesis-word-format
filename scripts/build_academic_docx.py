#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成中文学术格式 Word 文档（课程作业/论文通用排版）。

用法：
    python build_academic_docx.py --spec spec.json --out 输出.docx

spec.json 结构：
{
  "title": "论文题目",
  "author": "作者姓名　学号　第X次作业",        // 楷体小四，居中
  "abstract": "摘要文本（可选，置关键词前）",
  "keywords": "关键词1；关键词2（可选）",
  "sections": [                                 // 按顺序的内容块
    {"level": 1, "text": "1　引言"},            // level: 0=正文段 1=一级标题 2=二级标题
    {"level": 0, "text": "正文段落……"},
    {"table": 0},                               // 引用 tables 中下标为 0 的表格
    {"level": 2, "text": "2.1　小节"}
  ],
  "tables": [
    {
      "caption": "表1　标题",
      "headers": ["列1", "列2", "列3"],
      "rows": [["a", "b", "c"], ["d", "e", "f"]],
      "align": ["center", "left", "left"],      // 每列水平对齐（表头恒居中）
      "widths_cm": [2.8, 6.6, 6.6]              // 可选，不填则按内容均分
    }
  ],
  "references": ["文献1……", "文献2……"]          // 可选，自动生成 [1][2]… 编号列表
}

格式规则见同目录 ../references/format-spec.md。
依赖：python-docx（pip install python-docx）。
"""
import argparse
import copy
import json
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------- 基础工具 ----------------

def set_run(run, cn='宋体', en='Times New Roman', size=10.5, bold=False):
    run.font.name = en
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), cn)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def set_first_line_chars(p, chars=200, font_pt=10.5):
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    ind.set(qn('w:firstLineChars'), str(chars))
    ind.set(qn('w:firstLine'), str(int(chars / 100 * font_pt * 20)))


def mk(tag, attrs):
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn(k), v)
    return el


# ---------------- 文档初始化 ----------------

def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)   # A4
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Cm(2.5)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    normal.font.color.rgb = RGBColor(0, 0, 0)
    npf = normal.paragraph_format
    npf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    npf.line_spacing = 1.5
    npf.space_before = Pt(0)
    npf.space_after = Pt(0)

    for name, size, before, after in (('Heading 1', 12, 12, 6), ('Heading 2', 10.5, 6, 3)):
        st = doc.styles[name]
        st.font.name = 'Times New Roman'
        st.element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
    return doc


# ---------------- 段落构建 ----------------

def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    set_run(p.add_run(text), '黑体', 'Times New Roman', 18, bold=True)
    return p


def add_author(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    set_run(p.add_run(text), '楷体', 'Times New Roman', 12)
    return p


def add_abstract(doc, label, content):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    set_run(p.add_run(label), '黑体', 'Times New Roman', 9)
    set_run(p.add_run(content), '仿宋', 'Times New Roman', 9)
    return p


def add_heading(doc, level, text):
    style = {1: 'Heading 1', 2: 'Heading 2'}.get(level, 'Heading 1')
    size = {1: 12, 2: 10.5}.get(level, 12)
    p = doc.add_paragraph(style=style)
    set_run(p.add_run(text), '黑体', 'Times New Roman', size, bold=True)
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    set_first_line_chars(p, 200, 10.5)
    set_run(p.add_run(text), '宋体', 'Times New Roman', 10.5)
    return p


# ---------------- 表格（三线表） ----------------

def add_table(doc, spec):
    caption = spec.get('caption', '')
    if caption:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        set_run(p.add_run(caption), '黑体', 'Times New Roman', 9)

    headers = spec['headers']
    rows = spec.get('rows', [])
    align = spec.get('align') or ['center'] + ['left'] * (len(headers) - 1)
    ncols = len(headers)
    nrows = 1 + len(rows)
    tbl = doc.add_table(rows=nrows, cols=ncols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    tblPr = tbl._tbl.tblPr
    for child in list(tblPr):
        tblPr.remove(child)
    usable = 16.0  # A4 - 2*2.5cm 边距 = 16cm
    widths_cm = spec.get('widths_cm')
    if not widths_cm or len(widths_cm) != ncols:
        widths_cm = [usable / ncols] * ncols
    total = sum(widths_cm)
    widths_cm = [w * usable / total for w in widths_cm]
    widths_tw = [int(w / 2.54 * 1440) for w in widths_cm]

    tblPr.append(mk('w:tblW', {'w:w': str(int(usable / 2.54 * 1440)), 'w:type': 'dxa'}))
    tblPr.append(mk('w:jc', {'w:val': 'center'}))
    borders = OxmlElement('w:tblBorders')
    for tag, sz in (('top', 8), ('bottom', 8)):
        borders.append(mk('w:' + tag, {'w:val': 'single', 'w:sz': str(sz), 'w:space': '0', 'w:color': '000000'}))
    for tag in ('left', 'right', 'insideH', 'insideV'):
        borders.append(mk('w:' + tag, {'w:val': 'none', 'w:sz': '0', 'w:space': '0'}))
    tblPr.append(borders)
    tblPr.append(mk('w:tblLayout', {'w:type': 'fixed'}))
    cellmar = OxmlElement('w:tblCellMar')
    for tag, v in (('top', 0), ('left', 108), ('bottom', 0), ('right', 108)):
        cellmar.append(mk('w:' + tag, {'w:w': str(v), 'w:type': 'dxa'}))
    tblPr.append(cellmar)

    grid = tbl._tbl.find(qn('w:tblGrid'))
    for gc, w in zip(grid.findall(qn('w:gridCol')), widths_tw):
        gc.set(qn('w:w'), str(w))

    all_rows = [headers] + rows
    for i, row_data in enumerate(all_rows):
        row = tbl.rows[i]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
        if i == 0:
            trPr.append(OxmlElement('w:tblHeader'))
        for j, text in enumerate(row_data):
            cell = row.cells[j]
            cell.width = Cm(widths_cm[j])
            if i == 0:
                tcPr = cell._tc.get_or_add_tcPr()
                tcb = OxmlElement('w:tcBorders')
                tcb.append(mk('w:bottom', {'w:val': 'single', 'w:sz': '4', 'w:space': '0', 'w:color': '000000'}))
                tcPr.append(tcb)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else (
                WD_ALIGN_PARAGRAPH.CENTER if align[j] == 'center' else
                (WD_ALIGN_PARAGRAPH.RIGHT if align[j] == 'right' else WD_ALIGN_PARAGRAPH.LEFT))
            pf = p.paragraph_format
            pf.line_spacing = 1.0
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.first_line_indent = Pt(0)
            set_run(p.add_run(str(text)), '宋体', 'Times New Roman', 9, bold=(i == 0))
    return tbl


# ---------------- 参考文献（原生 [n] 编号） ----------------

def _prepare_reference_numbering(doc):
    """确保 numbering 部件存在，并把 List Number 的 1 级编号改为 [%1] 格式。返回该 numId。"""
    p = doc.add_paragraph('__probe__', style='List Number')   # 触发 numbering 部件创建
    doc._body._body.remove(p._p)
    num_elm = doc.part.numbering_part.element
    # 找到 List Number 样式绑定的 numId
    styles_elm = doc.styles.element
    listnum_id = None
    for st in styles_elm.findall(qn('w:style')):
        if st.get(qn('w:styleId')) == 'ListNumber':
            numpr = st.find(qn('w:pPr') + '/' + qn('w:numPr'))
            if numpr is not None:
                nid = numpr.find(qn('w:numId'))
                if nid is not None:
                    listnum_id = nid.get(qn('w:val'))
            break
    if listnum_id is None:
        # 兜底：取编号部件中第一个 num
        num = num_elm.find(qn('w:num'))
        listnum_id = num.get(qn('w:numId'))
    # 找到 num -> abstractNumId
    abstract_id = None
    for num in num_elm.findall(qn('w:num')):
        if num.get(qn('w:numId')) == listnum_id:
            an = num.find(qn('w:abstractNumId'))
            if an is not None:
                abstract_id = an.get(qn('w:val'))
            break
    # 修改 abstractNum 第 1 级编号文本为 [%1]
    for an in num_elm.findall(qn('w:abstractNum')):
        if an.get(qn('w:abstractNumId')) == abstract_id:
            for lvl in an.findall(qn('w:lvl')):
                if lvl.get(qn('w:ilvl')) == '0':
                    for lvlText in lvl.findall(qn('w:lvlText')):
                        lvlText.set(qn('w:val'), '[%1]')
                    for suff in lvl.findall(qn('w:suff')):
                        suff.set(qn('w:val'), 'space')
            break
    return listnum_id


def add_references(doc, refs):
    if not refs:
        return
    add_heading(doc, 1, '参考文献')
    num_id = _prepare_reference_numbering(doc)
    for text in refs:
        p = doc.add_paragraph(style='List Number')
        pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.line_spacing = 1.0
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.left_indent = Pt(18)          # 9pt 字号下 2 字符
        pf.first_line_indent = Pt(-18)   # 悬挂缩进
        set_run(p.add_run(text), '宋体', 'Times New Roman', 9)
    # 确保该批段落绑定到同一 numId
    for p in doc.paragraphs[-len(refs):]:
        pPr = p._p.get_or_add_pPr()
        numPr = pPr.find(qn('w:numPr'))
        if numPr is None:
            numPr = OxmlElement('w:numPr')
            pPr.append(numPr)
        ilvl = numPr.find(qn('w:ilvl'))
        if ilvl is None:
            ilvl = OxmlElement('w:ilvl')
            ilvl.set(qn('w:val'), '0')
            numPr.append(ilvl)
        nid = numPr.find(qn('w:numId'))
        if nid is None:
            nid = OxmlElement('w:numId')
            numPr.append(nid)
        nid.set(qn('w:val'), num_id)


# ---------------- 页脚页码 ----------------

def add_page_number_footer(doc):
    footer_p = doc.sections[0].footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer_p.add_run()
    set_run(fr, '宋体', 'Times New Roman', 9)
    fc1 = OxmlElement('w:fldChar'); fc1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = 'PAGE'
    fc2 = OxmlElement('w:fldChar'); fc2.set(qn('w:fldCharType'), 'end')
    fr._r.append(fc1); fr._r.append(instr); fr._r.append(fc2)


# ---------------- 主流程 ----------------

def build(spec, out_path):
    doc = new_document()
    add_title(doc, spec['title'])
    if spec.get('author'):
        add_author(doc, spec['author'])
    if spec.get('abstract'):
        add_abstract(doc, '摘要：', spec['abstract'])
    if spec.get('keywords'):
        add_abstract(doc, '关键词：', spec['keywords'])

    for block in spec.get('sections', []):
        if 'table' in block:
            add_table(doc, spec['tables'][block['table']])
        else:
            level = block.get('level', 0)
            if level == 0:
                add_body(doc, block['text'])
            else:
                add_heading(doc, level, block['text'])

    add_references(doc, spec.get('references', []))
    add_page_number_footer(doc)

    cp = doc.core_properties
    cp.title = spec.get('title', '')
    cp.author = spec.get('author', '')
    doc.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(description='生成中文学术格式 Word 文档')
    ap.add_argument('--spec', required=True, help='JSON spec 文件路径')
    ap.add_argument('--out', required=True, help='输出 .docx 路径')
    args = ap.parse_args()
    with open(args.spec, encoding='utf-8') as f:
        spec = json.load(f)
    build(spec, args.out)
    print('saved:', args.out)


if __name__ == '__main__':
    main()
